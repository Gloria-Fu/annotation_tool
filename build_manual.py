from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/fushiqing/code/annotate_tool')
OUT = ROOT / 'docs' / '标注管理平台操作手册.docx'
IMG = ROOT / 'docs' / 'manual_images'
IMG.mkdir(exist_ok=True)
FONT = '/System/Library/Fonts/Hiragino Sans GB.ttc'


def font(size):
    return ImageFont.truetype(FONT, size=size)


def draw_shell(title, subtitle, menu, cards, rows, filename):
    im = Image.new('RGB', (1400, 760), '#f5f7fa')
    d = ImageDraw.Draw(im)
    sidebar = 270
    d.rectangle((0, 0, sidebar, 760), fill='#152a3d')
    d.text((30, 25), 'OLA', font=font(34), fill='white')
    d.text((30, 70), '标注管理平台', font=font(23), fill='white')
    d.text((30, 106), 'Annotation Workspace', font=font(16), fill='#9db1c2')
    y = 165
    for label in menu:
        active = label == title or (title == '任务包' and label in ('标注任务包', '审核任务包'))
        d.rounded_rectangle((18, y, sidebar - 18, y + 48), radius=6, fill='#254661' if active else '#152a3d')
        d.text((36, y + 11), label, font=font(19), fill='white')
        y += 57
    d.rectangle((sidebar, 0, 1400, 76), fill='white')
    d.text((sidebar + 38, 24), '示例机器人项目', font=font(19), fill='#25394d')
    d.text((1170, 25), '研发管理员', font=font(17), fill='#536577')
    d.text((1320, 25), '退出', font=font(17), fill='#536577')
    d.text((sidebar + 40, 112), title, font=font(34), fill='#172333')
    d.text((sidebar + 40, 158), subtitle, font=font(17), fill='#657386')
    x = sidebar + 40
    for label, value, color in cards:
        d.rounded_rectangle((x, 220, x + 200, 330), radius=8, fill='white', outline='#d9e0e7', width=2)
        d.text((x + 18, 239), label, font=font(16), fill='#68788a')
        d.text((x + 18, 278), value, font=font(27), fill=color)
        x += 225
    d.rectangle((sidebar + 40, 370, 1355, 680), fill='white', outline='#d9e0e7', width=2)
    d.text((sidebar + 65, 395), '列表与操作', font=font(21), fill='#172333')
    d.line((sidebar + 65, 435, 1325, 435), fill='#d9e0e7', width=2)
    for i, row in enumerate(rows):
        yy = 455 + i * 52
        if i % 2 == 1:
            d.rectangle((sidebar + 55, yy - 8, 1335, yy + 35), fill='#f8fafc')
        d.text((sidebar + 75, yy), row, font=font(18), fill='#394a5a')
    im.save(IMG / filename)


draw_shell('进度看板', '项目任务状态与人员完成情况',
           ['进度看板', '人员统计', '项目管理', '账号管理', '数据导入', '任务包', '质量抽检'],
           [('任务总数', '12', '#1877b8'), ('待标注', '6', '#1877b8'), ('审核队列', '4', '#d97706'), ('完成率', '33%', '#16805b'), ('有效视频', '3 分钟', '#4958a8')],
           ['标注员 01    完成条目 4    有效时长 2 分钟', '审核员 01    完成条目 4    有效时长 1 分钟', '示例数据仅用于演示页面布局'], 'dashboard.png')
draw_shell('任务包', '每个任务条目对应一个 LeRobot episode',
           ['标注任务包', '我的标注', '我的工作量'],
           [('任务包', '1 个', '#1877b8'), ('总条目', '12', '#1877b8'), ('已标注', '6', '#16805b'), ('已审核', '4', '#16805b')],
           ['第一批抓取任务    已发布    顺序    总数 12 · 已领取 8 · 已标注 6 · 已审核 4', '操作：领取标注   管理条目', '提示：领取后任务会出现在“我的标注”'], 'packages.png')
draw_shell('我的标注', '领取任务、编辑片段并提交审核',
           ['标注任务包', '我的标注', '我的工作量'],
           [('待处理', '2', '#d97706'), ('已提交', '4', '#16805b'), ('草稿', '1', '#1877b8')],
           ['Episode 1    标注中    打开工作台', 'Episode 2    待审核    已提交', '操作：继续标注   查看历史'], 'workbench.png')
draw_shell('我的审核', '检查标注版本并作出审核决定',
           ['审核任务包', '我的工作量', '我的审核'],
           [('待审核', '2', '#d97706'), ('已通过', '4', '#16805b'), ('已退回', '1', '#c2410c')],
           ['Episode 3    审核中    打开工作台', 'Episode 4    已完成    查看历史', '操作：通过   退回修改'], 'review.png')


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders')
        tblPr.append(b)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        elem = b.find(qn('w:' + edge))
        if elem is None:
            elem = OxmlElement('w:' + edge); b.append(elem)
        elem.set(qn('w:val'), 'single'); elem.set(qn('w:sz'), '6'); elem.set(qn('w:space'), '0'); elem.set(qn('w:color'), color)


def add_img(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(path), width=Inches(6.55))
    c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    c.runs[0].italic = True; c.runs[0].font.size = Pt(9); c.runs[0].font.color.rgb = RGBColor(90, 100, 110)


doc = Document(); sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
styles = doc.styles
styles['Normal'].font.name = 'Hiragino Sans GB'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
styles['Normal'].font.size = Pt(10.5); styles['Normal'].paragraph_format.space_after = Pt(5); styles['Normal'].paragraph_format.line_spacing = 1.15
for name in ('Title', 'Heading 1', 'Heading 2'):
    styles[name].font.name = 'Hiragino Sans GB'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB'); styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 1'].paragraph_format.space_before = Pt(13); styles['Heading 1'].paragraph_format.space_after = Pt(5)
styles['Heading 2'].paragraph_format.space_before = Pt(8); styles['Heading 2'].paragraph_format.space_after = Pt(3)

doc.add_heading('标注管理平台操作手册', 0)
p = doc.add_paragraph('外包团队快速上手指南'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size = Pt(16); p.runs[0].font.color.rgb = RGBColor(52, 73, 94)
p = doc.add_paragraph('适用角色：标注员、审核员、标注管理员、研发管理员    版本：v1.1'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size = Pt(10)
doc.add_paragraph('本手册用“先做什么、再看哪里、完成后如何确认”的方式说明平台操作。操作路径已在项目同版前端的本地真实运行实例中核对；由于远端生产环境无法登录，本文配图采用界面示意图，图中数据为示例数据。正式任务以项目管理员分配的项目和任务包为准。')

doc.add_heading('一 先记住这条工作链路', 1)
doc.add_paragraph('外包团队的日常工作可以按以下顺序完成：登录 → 选择项目 → 领取任务 → 打开工作台 → 播放并切分片段 → 填写精细字段 → 保存草稿 → 提交审核（标注员）或通过/退回（审核员）。')
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
for i, text in enumerate(['角色', '每天要做的事', '完成标志']):
    table.columns[i].width = Inches([1.3, 3.7, 1.5][i]); table.rows[0].cells[i].text = text; shade(table.rows[0].cells[i], '1F4E78')
    for run in table.rows[0].cells[i].paragraphs[0].runs: run.font.color.rgb = RGBColor(255, 255, 255); run.bold = True
for values in [('标注员', '领取标注任务，完成片段和精细字段', '点击提交审核'), ('审核员', '领取审核任务，核对边界、内容和字段', '点击通过或退回修改'), ('标注管理员', '分配任务、回收任务、抽检质量', '任务包进度持续更新'), ('研发管理员', '维护项目、账号、数据集和全局进度', '看板与报表可追溯')]:
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value; cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if len(table.rows) % 2 == 0: shade(cells[i], 'F7FAFC')
borders(table)

doc.add_heading('二 登录与页面导航', 1)
doc.add_paragraph('在浏览器打开项目提供的地址，输入用户名和密码后点击“登录”。首次登录若弹出“请修改密码”，先按要求完成修改，再继续操作。登录后，左侧是功能导航，顶部是项目选择器和当前用户。')
doc.add_paragraph('外包团队最常用的三个入口：')
for text in ['“标注任务包”：查看可领取的标注任务。', '“我的标注”：继续处理已领取的任务，查看草稿和已提交任务。', '“审核任务包 / 我的审核”：领取审核任务并查看审核历史。']:
    doc.add_paragraph(text, style='List Bullet')
add_img(doc, IMG / 'dashboard_real.png', '图 1 进度看板（本地真实运行页面截图；数据为演示数据）')

doc.add_heading('三 标注员操作', 1)
doc.add_heading('1 领取标注任务', 2)
doc.add_paragraph('进入“标注任务包”，确认顶部项目正确，找到状态为“已发布”的任务包，点击“领取标注”。领取成功后系统会打开任务工作台；也可以稍后在“我的标注”中继续处理。领取失败时，先检查任务包是否已被领完、账号是否属于该项目。')
add_img(doc, IMG / 'packages.png', '图 2 标注任务包列表与领取入口（界面示意图；操作已在本地真实页面核对，数据为示例）')
doc.add_heading('2 在工作台完成一条任务', 2)
for text in ['播放视频：用播放器确认动作发生的时间范围；没有视频时先联系项目管理员，不要凭猜测填写。', '选择片段：在时间轴点击目标片段。播放头位于片段内部时，点击“分段”可拆分；拖动片段边界可微调起止帧。', '填写精细字段：在片段编辑区选择 Skill，并填写项目要求的手、位置、物体、动作结果、失败原因和关键帧等字段。', '检查预览句：预览文本应能读懂“谁在什么位置，对什么物体做了什么动作，结果如何”。发现空字段或语义不通时，回到当前片段修改。']:
    doc.add_paragraph(text, style='List Number')
add_img(doc, IMG / 'workbench.png', '图 3 我的标注与工作台入口（界面示意图；操作已在本地真实页面核对，数据为示例）')
doc.add_heading('3 保存与提交', 2)
doc.add_paragraph('平台支持自动保存，但在切换片段、离开页面或长时间操作前，建议点击“保存草稿”确认提示。所有片段的必填字段完成后，“提交审核”按钮才会可用。提交后任务进入“待审核”，标注员不能继续按草稿方式修改；如果审核退回，任务会回到待处理列表，修改后再次提交。')
doc.add_paragraph('常用快捷操作：空格键在当前片段内分段；“撤销 / 重做”用于回退最近一次边界或字段修改。误操作后先撤销，再继续编辑。')

doc.add_heading('四 审核员操作', 1)
doc.add_heading('1 领取审核任务', 2)
doc.add_paragraph('进入“审核任务包”，选择已发布任务包并点击“领取审核”。系统会弹出领取方式，可选择“顺序抽检”或“随机抽检”。领取成功后在“我的审核”打开任务。')
doc.add_heading('2 按四项检查', 2)
for text in ['时间边界：片段起止帧是否覆盖完整动作，是否包含多余等待。', '内容字段：Skill、物体和动作结果是否与视频一致。', '精细字段：手、位置、接触点、关键帧等必填项是否完整。', '语义预览：标注句是否通顺，是否能让其他人复现动作。']:
    doc.add_paragraph(text, style='List Bullet')
add_img(doc, IMG / 'review.png', '图 4 我的审核与审核决定入口（界面示意图；操作已在本地真实页面核对，数据为示例）')
doc.add_heading('3 通过或退回', 2)
doc.add_paragraph('确认无误后点击“通过”。发现问题时点击“退回修改”，在意见中写清“哪一个片段、哪一个字段、需要怎么改”，例如“片段 2 的结束帧多了约 10 帧，请拖回到物体离开桌面的时刻”。审核修改模式下可直接修正内容，但仍应保留清晰的审核意见。')

doc.add_heading('五 管理员操作概览', 1)
doc.add_paragraph('研发管理员和标注管理员负责把任务准备好并持续跟踪。外包团队通常只需关注任务包和自己的工作页面，但了解下面的顺序有助于定位问题：')
for index, text in enumerate(['项目管理：创建项目并确认项目名称。', '账号管理 / 人员管理：创建账号、分配角色和项目成员关系。', '数据导入：登记数据集路径，等待状态变为“就绪”。', '任务包：从就绪数据集创建任务包，设置任务数量和领取方式，发布后才可领取。', '任务条目：按标注阶段或审核阶段批量指派，也可回收已领取条目。', '质量抽检：从已审核任务生成按比例、按数量或全部抽检的批次，并记录通过/退回及原因。', '进度看板 / 人员统计：查看总量、待标注、审核队列、完成率和个人贡献。'], 1):
    doc.add_paragraph(f'{index}. {text}')

doc.add_heading('六 状态和异常处理', 1)
table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, text in enumerate(['页面状态', '你应该怎么做']):
    table.rows[0].cells[i].text = text; shade(table.rows[0].cells[i], '1F4E78')
    for run in table.rows[0].cells[i].paragraphs[0].runs: run.font.color.rgb = RGBColor(255, 255, 255); run.bold = True
for a, b in [('可领取 / 已发布', '可以领取对应阶段的任务。'), ('标注中', '标注员已领取，继续处理并保存。'), ('待审核', '标注已提交，等待审核员处理。'), ('审核中', '审核员已领取，完成核对后通过或退回。'), ('退回修改', '阅读审核意见，修改后重新提交。'), ('已完成', '审核通过；不要重复领取或修改。')]:
    cells = table.add_row().cells; cells[0].text = a; cells[1].text = b
    if len(table.rows) % 2 == 0: shade(cells[0], 'F7FAFC'); shade(cells[1], 'F7FAFC')
borders(table)
doc.add_paragraph('遇到“无法领取”：确认项目、角色、任务包状态和是否已有领取记录。遇到“无法提交”：逐段检查是否选择 Skill、是否填写必填字段，先保存草稿再重试。遇到“页面报错或数据不更新”：刷新页面并重新选择项目；仍未恢复时，把用户名、项目名、任务条目 ID、页面提示和发生时间发给项目管理员。')
doc.add_paragraph('数据安全提醒：不要修改原始数据集的元数据、Parquet 文件或视频；不要把账号密码写入截图、群聊或工单。')

doc.add_page_break()
doc.add_heading('附录 外包团队交付前自检清单', 1)
for text in ['当前项目与任务包一致。', '每个片段起止帧与动作实际发生范围一致。', '所有必填精细字段已填写，预览句通顺。', '已保存草稿，提交后状态变为“待审核”。', '退回任务已按审核意见逐条修改并再次提交。']:
    doc.add_paragraph('□ ' + text)

doc.save(OUT)
print(OUT)
